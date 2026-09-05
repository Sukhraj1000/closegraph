"""Synthetic schema-bound export tests, including hostile spreadsheet values."""
import csv
from datetime import date
from io import BytesIO, StringIO
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
import pytest

from closegraph.collections.export import export_table


def table(rows=None):
    return {"table_id": "reviewed-data", "source_id": "original-source", "columns": [
        {"key": "identifier", "label": "Account"}, {"key": "amount", "label": "Amount"}],
        "rows": [{"row_id": f"row-{index}", "values": {"identifier": key, "amount": amount},
                  "locators": {"amount": {"kind": "xlsx", "sheet": "Source", "cell": f"D{index + 5}"}}}
                 for index, (key, amount) in enumerate(rows or [("00123", "0.30")])]}


def workbook_bytes(book):
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def test_csv_roundtrip_preserves_identifiers_delimiters_and_source_manifest():
    result = export_table(table([("00123", "0.30"), ('Comma, quote" and\nnewline', "42.17")]), "csv")
    rows = list(csv.reader(StringIO(result["content"].decode("utf-8-sig"))))
    assert rows == [["Account", "Amount"], ["00123", "0.30"], ['Comma, quote" and\nnewline', "42.17"]]
    assert result["verification"]["readback_passed"]
    assert result["verification"]["source_manifest"]["cells"][0]["columns"]["amount"][0]["locator"]["cell"] == "D5"


@pytest.mark.parametrize("text", ['=HYPERLINK("https://invalid")', "+1+2", "-1+2", "@SUM(A1)", "  =1+1", "\t=1+1", "\n=1+1"])
def test_csv_formula_prefixes_are_escaped_and_recorded(text):
    result = export_table(table([(text, "12")]))
    rows = list(csv.reader(StringIO(result["content"].decode("utf-8-sig"))))
    assert rows[1][0] == "'" + text
    assert result["verification"]["escaped_cells"] == [{"row_id": "row-0", "column": "identifier"}]


def test_csv_header_formula_is_also_escaped():
    source = table()
    source["columns"][0]["label"] = "=1+1"
    result = export_table(source)
    assert list(csv.reader(StringIO(result["content"].decode("utf-8-sig"))))[0][0] == "'=1+1"
    assert result["verification"]["escaped_cells"][0]["row"] == "header"


def test_xlsx_defaults_literal_text_preserving_precision_and_formula_like_values():
    source = table([("=1+1", "12345678901234567890.123456789"), ("00007", "-12.00")])
    result = export_table(source, "xlsx")
    book = load_workbook(BytesIO(result["content"]), data_only=False)
    assert book["Export"]["A2"].value == "=1+1"
    assert book["Export"]["A2"].data_type == "s"
    assert book["Export"]["B2"].value == "12345678901234567890.123456789"
    assert book["Export"]["A3"].value == "00007"
    assert result["verification"]["readback_passed"]


def test_explicit_output_mapping_order_headers_and_decimal_type():
    binding = {"sheet": "Custom report", "start_row": 4, "header_row": 3,
               "columns": {"identifier": "D", "amount": "B"}, "types": {"amount": "decimal"},
               "labels": {"amount": "Net amount"}}
    result = export_table(table(), "xlsx", bindings=binding)
    sheet = load_workbook(BytesIO(result["content"]))["Custom report"]
    assert sheet["B3"].value == "Net amount"
    assert sheet["B4"].value == 0.3
    assert sheet["D4"].value == "00123"
    assert sheet["D4"].data_type == "s"
    csv_result = export_table(table(), "csv", bindings=binding)
    assert list(csv.reader(StringIO(csv_result["content"].decode("utf-8-sig")))) == [["Net amount", "Account"], ["0.30", "00123"]]


def test_explicit_date_binding_roundtrips_iso_dates():
    result = export_table(table([("2025-12-31", "1")]), "xlsx", bindings={"types": {"identifier": "date"}})
    actual = load_workbook(BytesIO(result["content"]))["Export"]["A2"].value
    assert actual.date() == date(2025, 12, 31)


@pytest.mark.parametrize("value,kind", [("1234567890123456", "decimal"), ("1.5", "integer"), ("000invalid", "decimal"), ("31/12/2025", "date")])
def test_invalid_or_lossy_explicit_types_fail_without_silent_coercion(value, kind):
    with pytest.raises(ValueError):
        export_table(table([(value, "1")]), "xlsx", bindings={"types": {"identifier": kind}})


def test_template_preserves_unmapped_values_styles_and_sheets_with_variable_rows():
    book = Workbook()
    sheet = book.active
    sheet.title = "Loader"
    sheet["A1"] = "Template title"
    sheet.merge_cells("A1:D1")
    sheet["A3"] = "Account"
    sheet["C3"] = "Amount"
    sheet["F12"] = "Keep this footer"
    sheet["F12"].font = Font(bold=True, color="FF112233")
    sheet["A4"].fill = PatternFill("solid", fgColor="FFFFFF00")
    sheet.column_dimensions["A"].width = 25
    sheet.freeze_panes = "A4"
    other = book.create_sheet("Reference")
    other["B2"] = "Unmapped reference value"
    content = workbook_bytes(book)
    result = export_table(table([("001", "12"), ("002", "13"), ("003", "14")]), "xlsx", template_bytes=content,
                          bindings={"sheet": "Loader", "start_row": 4, "header_row": 3, "columns": {"identifier": "A", "amount": "C"}})
    actual = load_workbook(BytesIO(result["content"]))
    assert actual.sheetnames == ["Loader", "Reference"]
    assert actual["Loader"]["A6"].value == "003"
    assert actual["Loader"]["C6"].value == "14"
    assert actual["Loader"]["F12"].value == "Keep this footer"
    assert actual["Loader"]["F12"].font.bold
    assert actual["Loader"]["A4"].fill.fgColor.rgb == "FFFFFF00"
    assert actual["Reference"]["B2"].value == "Unmapped reference value"
    assert result["verification"]["template_preserved"]


def test_template_does_not_overwrite_nonblank_data_or_move_footer():
    book = Workbook()
    book.active.title = "Loader"
    book.active["A3"] = "Do not replace"
    with pytest.raises(ValueError, match="overwrite"):
        export_table(table([("new", "1"), ("second", "2")]), "xlsx", template_bytes=workbook_bytes(book),
                     bindings={"sheet": "Loader", "start_row": 2, "columns": {"identifier": "A"}})


def test_template_requires_bindings_and_rejects_merged_target():
    book = Workbook()
    book.active.title = "Loader"
    book.active.merge_cells("A2:B2")
    template = workbook_bytes(book)
    with pytest.raises(ValueError, match="explicit"):
        export_table(table(), "xlsx", template_bytes=template)
    with pytest.raises(ValueError, match="merged"):
        export_table(table(), "xlsx", template_bytes=template, bindings={"sheet": "Loader", "start_row": 2, "columns": {"identifier": "A"}})


def test_unsupported_formula_template_is_explicitly_rejected():
    book = Workbook()
    book.active["D10"] = "=SUM(A1:A9)"
    with pytest.raises(ValueError, match="formulas"):
        export_table(table(), "xlsx", template_bytes=workbook_bytes(book), bindings={"sheet": "Sheet", "start_row": 2, "columns": {"identifier": "A"}})


def test_unsupported_template_package_parts_are_rejected_before_loading():
    source = workbook_bytes(Workbook())
    stream = BytesIO()
    with zipfile.ZipFile(BytesIO(source)) as old, zipfile.ZipFile(stream, "w") as new:
        for entry in old.infolist():
            new.writestr(entry.filename, old.read(entry.filename))
        new.writestr("xl/vbaProject.bin", b"unsupported")
    with pytest.raises(ValueError, match="Unsupported template feature"):
        export_table(table(), "xlsx", template_bytes=stream.getvalue(), bindings={"sheet": "Sheet", "start_row": 2, "columns": {"identifier": "A"}})


@pytest.mark.parametrize("bindings", [
    {"columns": {"unknown": "A"}}, {"columns": {"identifier": "A", "amount": "A"}},
    {"start_row": 0}, {"header_row": 2}, {"columns": {"identifier": "XFE"}},
    {"types": {"identifier": "formula"}}, {"unexpected": True},
])
def test_invalid_schema_bindings_fail_closed(bindings):
    with pytest.raises(ValueError):
        export_table(table(), "xlsx", bindings=bindings)


def test_xlsx_length_and_control_characters_do_not_silently_truncate():
    with pytest.raises(ValueError, match="character limit"):
        export_table(table([("a" * 32768, "1")]), "xlsx")
    with pytest.raises(ValueError, match="control"):
        export_table(table([("a\x00b", "1")]), "csv")
