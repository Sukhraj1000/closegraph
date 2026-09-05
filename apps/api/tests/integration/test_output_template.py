from io import BytesIO
from copy import copy
from decimal import Decimal
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from closegraph.outputs.workbook import TemplateContract, generate_workbook, verify_roundtrip, digest
from closegraph.services.errors import DomainError


def template():
    w=Workbook(); s=w.active; s.title="Pack"; s["A1"]="SYNTHETIC reporting pack"; s["A1"].font=Font(bold=True)
    s["A2"]="Net assets"; s["B2"]=0; s["B2"].number_format="#,##0.00"; s.column_dimensions["A"].width=30; s.freeze_panes="B2"
    b=BytesIO(); w.save(b); data=b.getvalue()
    return data, TemplateContract(digest(data),(("total","Pack","B2"),))


def test_real_xlsx_roundtrip_preserves_values_layout_styles():
    original,contract=template()
    data=generate_workbook(original,contract,{"total":"990.00"})
    w=load_workbook(BytesIO(data)); old=load_workbook(BytesIO(original))
    assert w["Pack"]["B2"].value == 990
    assert w["Pack"]["A1"].value == old["Pack"]["A1"].value
    assert copy(w["Pack"]["A1"].font) == copy(old["Pack"]["A1"].font)
    assert w["Pack"].column_dimensions["A"].width == 30
    assert w["Pack"].freeze_panes == "B2"
    assert verify_roundtrip(original,data,contract,{"total":"990.00"}) == digest(data)
    assert load_workbook(BytesIO(original))["Pack"]["B2"].value == 0


@pytest.mark.parametrize("kind", ["formula","unmapped","precision","unknown-binding","hash"])
def test_template_denials(kind):
    data,c=template(); values={"total":"990.00"}
    if kind == "formula":
        w=load_workbook(BytesIO(data)); w["Pack"]["A3"]="=1+1"; b=BytesIO(); w.save(b); data=b.getvalue(); c=TemplateContract(digest(data),c.bindings)
    if kind == "precision": values={"total":"1234567890123456"}
    if kind == "unknown-binding": values["extra"]="5"
    if kind == "hash": c=TemplateContract("f"*64,c.bindings)
    with pytest.raises(DomainError):
        if kind == "unmapped":
            out=generate_workbook(data,c,values); w=load_workbook(BytesIO(out)); w["Pack"]["A1"]="tampered"; b=BytesIO(); w.save(b)
            verify_roundtrip(data,b.getvalue(),c,values)
        else: generate_workbook(data,c,values)
