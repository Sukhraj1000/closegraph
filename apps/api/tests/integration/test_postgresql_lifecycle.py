"""Real PostgreSQL acceptance tests; never substitute SQLite or memory persistence."""
import base64
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
import os
from pathlib import Path
from uuid import uuid4

import pytest
from alembic import command
from alembic.config import Config
from openpyxl import Workbook, load_workbook
from sqlalchemy import text, select

from closegraph.api.auth import DevAccount, LocalAuth
from closegraph.api.ports import Actor, DomainConflict, DomainNotFound
from closegraph.api.services import PostgreSQLPackServices
from closegraph.contracts import Scope, CorrectionRequest, VersionAction, ReviewRequest, UploadRequest, ResolutionRequest
from closegraph.services.repository import ScopedBlobs
from closegraph.storage.blobs import LocalBlobStore
from closegraph.storage.database import make_engine, session_factory
from closegraph.storage.models import (
    PackRevisionRow, FactVersionRow, DocumentVersionRow, ProcessingResultRow, ReviewDecisionRow, PublicationRow,
)

SCOPE = Scope(tenant_id="tenant",fund_id="fund",pack_id="pack")
PREPARER = Actor(actor_id="preparer",role="PREPARER")
REVIEWER = Actor(actor_id="reviewer",role="REVIEWER")


class Evaluator:
    required_check_ids = ("fee",)
    code_version = "postgres-test-v1"

    def __call__(self, state, blobs):
        facts = {f["fact_id"]:f for f in state["facts"]}
        source_id = facts["fee"].get("correction_source_id") or facts["fee"]["source"]["source_id"]
        expected = Decimal(blobs.get(state["evidence"][source_id]["content_hash"]).decode().strip())
        actual = Decimal(facts["fee"]["value_decimal"])
        capital = Decimal(facts["capital"]["value_decimal"])
        return {"checks":[{"id":"fee","required":True,"status":"PASS" if actual==expected else "FAIL",
            "label":"Fee ties to source", "difference":format(actual-expected,"f"),"tolerance":"0",
            "detail":"Source-backed fee comparison","operands":[{"fact_id":"fee","actual":str(actual),"expected":str(expected)}]}],
            "values":{"capital":str(capital),"fee":str(actual),"total":format(capital-actual,".2f")}}


@pytest.fixture
def database():
    url = os.environ.get("CLOSEGRAPH_TEST_DATABASE_URL")
    if not url:
        pytest.skip("CLOSEGRAPH_TEST_DATABASE_URL is required for actual PostgreSQL tests")
    admin = make_engine(url)
    schema = "test_" + uuid4().hex
    with admin.begin() as connection:
        connection.execute(text("CREATE SCHEMA " + schema))
    engine = make_engine(url,connect_args={"options":"-csearch_path="+schema})
    config = Config(str(Path(__file__).parents[2]/"alembic.ini"))
    with engine.begin() as connection:
        config.attributes["connection"] = connection
        command.upgrade(config,"head")
    yield session_factory(engine)
    engine.dispose()
    with admin.begin() as connection:
        connection.execute(text("DROP SCHEMA "+schema+" CASCADE"))
    admin.dispose()


@pytest.fixture
def application(database,tmp_path):
    auth = LocalAuth(accounts=(
        DevAccount.create("preparer","preparer-test-password","PREPARER",(SCOPE,)),
        DevAccount.create("reviewer","reviewer-test-password","REVIEWER",(SCOPE,)),
    ))
    blobs = LocalBlobStore(tmp_path/"blobs")
    service = PostgreSQLPackServices(database,blobs,auth,Evaluator())
    scoped = ScopedBlobs(blobs,SCOPE)
    capital_hash,fee_hash = scoped.put(b"1000.00"),scoped.put(b"10.00")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pack"
    sheet.append(["Synthetic statement","Amount"])
    sheet.append(["Capital",1000])
    sheet.append(["Fee",20])
    sheet.append(["Net",980])
    sheet.column_dimensions["A"].width = 30
    output = BytesIO();workbook.save(output)
    template_hash = scoped.put(output.getvalue())
    evidence = {
        name:{"source_version":1,"content_hash":identity,"locator":{"kind":"csv","row":1,"column":"value"},
              "fact_ids":[fact],"entity_id":"entity","currency":"GBP","period":"2026-Q2"}
        for name,identity,fact in (("capital-source",capital_hash,"capital"),("fee-source",fee_hash,"fee"))
    }
    state = {"title":"Synthetic pack","facts":[
        {"fact_id":fact,"fact_version":1,"metric":fact,"value_decimal":value,"entity_id":"entity",
         "currency":"GBP","period":"2026-Q2","raw_value":value,"raw_scale":"1","source":{"source_id":source}}
        for fact,value,source in (("capital","1000.00","capital-source"),("fee","20.00","fee-source"))],
        "evidence":evidence,"source_versions":{name:1 for name in evidence},"rule_version":"source-fee-v1",
        "mapping_version":"template-v1","policy_version":"native-single-parser-v1",
        "dependency_edges":[["fee","output"],["capital","output"]],"dependency_coverage":True,
        "template":{"template_hash":template_hash,"version":"value-only-v1","bindings":[
            ["capital","Pack","B2"],["fee","Pack","B3"],["total","Pack","B4"]]},
        "observations":[{"primary":{},"secondary":None,"availability":"NOT_RUN"}],"native":True,
        "checks":[{"id":"fee","required":True,"status":"FAIL","detail":"20 differs from 10"}],
        "values":{},"history":[],"freshness":"STALE","execution_status":"COMPLETED",
        "routing_status":"BLOCKED","review_status":"PENDING","contributors":["preparer"],"publications":[]}
    service.create_pack(SCOPE,PREPARER,state)
    return service


def run_pending(service,version,*,suffix=""):
    pending=service.recompute(SCOPE,PREPARER,VersionAction(expected_version=version))
    request=service.pending_requests()[-1]
    loaded=service.load_processing(SCOPE,request["request_id"])
    assert service.claim_processing(SCOPE,request["request_id"],"run"+suffix)
    result=service.evaluator(loaded["state"],ScopedBlobs(service.blob_store,SCOPE))
    complete=service.complete_processing(SCOPE,request["request_id"],result,"run"+suffix)
    assert complete["applied"]
    return complete["state"]


def repair(service):
    state=service.get_pack(SCOPE,PREPARER)
    state=service.correct(SCOPE,PREPARER,CorrectionRequest(expected_version=state.version,fact_id="fee",
        value_decimal="10.00",reason="Read the actual scoped fee source",source_id="fee-source"))
    return run_pending(service,state.version)


def test_actual_postgres_failed_then_successful_correction_download_and_manifest(application,database):
    service=application
    bad=service.correct(SCOPE,PREPARER,CorrectionRequest(expected_version=1,fact_id="fee",
        value_decimal="9.00",reason="First attempt",source_id="fee-source"))
    bad=run_pending(service,bad.version,suffix="bad")
    assert bad.execution_status=="COMPLETED" and bad.routing_status=="BLOCKED" and bad.candidate is None
    with pytest.raises(DomainConflict):
        service.review(SCOPE,REVIEWER,ReviewRequest(expected_version=bad.version,note="Cannot waive",attested=True))
    fixed=repair(service)
    assert fixed.routing_status=="READY_FOR_QUICK_REVIEW" and fixed.review_status=="PENDING", fixed.model_dump()
    preview=service.download(SCOPE,PREPARER,fixed.candidate.artifact_id)
    assert "UNAPPROVED-PREVIEW" in preview.filename
    approved=service.review(SCOPE,REVIEWER,ReviewRequest(expected_version=fixed.version,note="Compared sources",attested=True))
    released=service.publish(SCOPE,REVIEWER,VersionAction(expected_version=approved.version))
    duplicate=service.publish(SCOPE,REVIEWER,VersionAction(expected_version=approved.version))
    assert duplicate.version==released.version
    output=service.download(SCOPE,REVIEWER,released.candidate.artifact_id)
    assert load_workbook(BytesIO(output.content)).active["B4"].value==990
    manifest=service.download(SCOPE,REVIEWER,released.publication["manifest_id"])
    import json
    decoded=json.loads(manifest.content)
    assert decoded["artifact_sha256"]==released.candidate.artifact_id
    assert decoded["review"]["actor"]=="reviewer" and decoded["sources"]
    with database() as session:
        facts=session.scalars(select(FactVersionRow).where(FactVersionRow.fact_id=="fee").order_by(FactVersionRow.version)).all()
        assert [row.value_decimal for row in facts]==[Decimal("20"),Decimal("9"),Decimal("10")]
        assert len(session.scalars(select(ReviewDecisionRow)).all())==1
        assert len(session.scalars(select(PublicationRow)).all())==1


def test_row_lock_cas_and_historical_immutability(application,database):
    command=CorrectionRequest(expected_version=1,fact_id="fee",value_decimal="10",reason="source",source_id="fee-source")
    application.correct(SCOPE,PREPARER,command)
    with pytest.raises(DomainConflict):
        application.correct(SCOPE,PREPARER,command)
    with database.begin() as session:
        with pytest.raises(Exception,match="append-only"):
            session.execute(text("UPDATE pack_revisions SET prepared_by='attacker'"))


def test_upload_receipt_idempotency_conflict_and_original_bytes(application,database):
    command=UploadRequest(expected_version=1,idempotency_key="upload-1",source_id="extra",filename="extra.csv",
        media_type="text/csv",content_base64=base64.b64encode(b"value\n1\n").decode())
    state=application.upload(SCOPE,PREPARER,command)
    assert state.execution_status=="PENDING"
    again=application.upload(SCOPE,PREPARER,command)
    assert state.version==again.version
    changed=command.model_copy(update={"content_base64":base64.b64encode(b"value\n2\n").decode()})
    with pytest.raises(DomainConflict):
        application.upload(SCOPE,PREPARER,changed)
    assert application.download_evidence(SCOPE,PREPARER,"extra").content==b"value\n1\n"
    with database() as session:
        assert len(session.scalars(select(DocumentVersionRow).where(DocumentVersionRow.source_id=="extra")).all())==1


def test_restart_and_late_processing_do_not_overwrite_newer_state(application,database):
    pending=application.recompute(SCOPE,PREPARER,VersionAction(expected_version=1))
    request=application.pending_requests()[0]
    loaded=application.load_processing(SCOPE,request["request_id"])
    application.claim_processing(SCOPE,request["request_id"],"late-run")
    corrected=application.correct(SCOPE,PREPARER,CorrectionRequest(expected_version=pending.version,
        fact_id="fee",value_decimal="10",reason="Newer correction",source_id="fee-source"))
    restarted=PostgreSQLPackServices(database,application.blob_store,application.auth,Evaluator())
    result=restarted.evaluator(loaded["state"],ScopedBlobs(application.blob_store,SCOPE))
    completed=restarted.complete_processing(SCOPE,request["request_id"],result,"late-run")
    assert not completed["applied"] and completed["state"].version==corrected.version
    repeated=restarted.complete_processing(SCOPE,request["request_id"],result,"late-run")
    assert not repeated["applied"]
    with database() as session:
        assert len(session.scalars(select(ProcessingResultRow)).all())==1


def test_cross_tenant_fund_and_pack_access_has_no_evidence_leak(application):
    scopes=[Scope(tenant_id="other",fund_id="fund",pack_id="pack"),
            Scope(tenant_id="tenant",fund_id="other",pack_id="pack"),
            Scope(tenant_id="tenant",fund_id="fund",pack_id="other")]
    for scope in scopes:
        with pytest.raises(DomainNotFound):
            application.get_pack(scope,PREPARER)
        with pytest.raises(DomainNotFound):
            application.download_evidence(scope,PREPARER,"fee-source")


def test_correction_invalidates_approval_and_historical_release_remains_labelled(application):
    state=repair(application)
    state=application.review(SCOPE,REVIEWER,ReviewRequest(expected_version=state.version,note="Sources checked",attested=True))
    state=application.publish(SCOPE,REVIEWER,VersionAction(expected_version=state.version))
    artifact=state.candidate.artifact_id
    changed=application.correct(SCOPE,PREPARER,CorrectionRequest(expected_version=state.version,fact_id="fee",
        value_decimal="9",reason="New disputed source",source_id="fee-source"))
    assert changed.review_status=="PENDING" and changed.freshness=="STALE"
    with pytest.raises(DomainConflict):
        application.publish(SCOPE,REVIEWER,VersionAction(expected_version=changed.version))
    old=application.download(SCOPE,REVIEWER,artifact)
    assert "stale" in old.filename


def test_soft_resolution_is_separate_and_required_before_approval(application):
    pending=application.correct(SCOPE,PREPARER,CorrectionRequest(expected_version=1,fact_id="fee",
        value_decimal="10",reason="Actual source fee",source_id="fee-source"))
    pending=application.recompute(SCOPE,PREPARER,VersionAction(expected_version=pending.version))
    request=application.pending_requests()[0]
    loaded=application.load_processing(SCOPE,request["request_id"])
    application.claim_processing(SCOPE,request["request_id"],"soft")
    result=application.evaluator(loaded["state"],ScopedBlobs(application.blob_store,SCOPE))
    result["observations"]=[{"primary":{"metric":"fee","value_decimal":"10","entity_id":"entity",
        "period":"2026-Q2","currency":"GBP","raw_scale":"1"},"secondary":{"metric":"fee","value_decimal":"11",
        "entity_id":"entity","period":"2026-Q2","currency":"GBP","raw_scale":"1"}}]
    state=application.complete_processing(SCOPE,request["request_id"],result,"soft")["state"]
    assert state.routing_status=="NEEDS_REVIEW"
    with pytest.raises(DomainConflict,match="resolution"):
        application.review(SCOPE,REVIEWER,ReviewRequest(expected_version=state.version,note="Checked",attested=True))
    resolved=application.resolve(SCOPE,REVIEWER,ResolutionRequest(expected_version=state.version,
        reason="Compared the actual CSV source",source_ids=("fee-source",)))
    assert resolved.review_status=="PENDING" and resolved.resolution
    approved=application.review(SCOPE,REVIEWER,ReviewRequest(expected_version=resolved.version,note="Checked",attested=True))
    assert approved.review_status=="APPROVED"
    assert [e["action"] for e in approved.history][-2:]==["RESOLVED","APPROVED"]


def test_independence_applies_to_correction_contributors_even_with_reviewer_grant(application):
    state=repair(application)
    application.auth._accounts["preparer"]=DevAccount.create("preparer","new-preparer-password","REVIEWER",(SCOPE,))
    actor=Actor(actor_id="preparer",role="REVIEWER")
    with pytest.raises(DomainConflict,match="Independent"):
        application.review(SCOPE,actor,ReviewRequest(expected_version=state.version,note="My own correction",attested=True))


def test_simultaneous_publication_and_correction_serialize_at_current_snapshot(application):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    state=repair(application)
    state=application.review(SCOPE,REVIEWER,ReviewRequest(expected_version=state.version,note="Sources checked",attested=True))
    barrier=Barrier(2)
    def publish():
        barrier.wait()
        try:
            return application.publish(SCOPE,REVIEWER,VersionAction(expected_version=state.version))
        except DomainConflict:
            return None
    def correct():
        barrier.wait()
        try:
            return application.correct(SCOPE,PREPARER,CorrectionRequest(expected_version=state.version,
                fact_id="fee",value_decimal="9",reason="New relevant version",source_id="fee-source"))
        except DomainConflict:
            return None
    with ThreadPoolExecutor(max_workers=2) as executor:
        first,second=executor.submit(publish),executor.submit(correct)
        outcomes=[first.result(),second.result()]
    assert sum(outcome is not None for outcome in outcomes)==1
    current=application.get_pack(SCOPE,PREPARER)
    assert current.version==state.version+1
    if current.candidate is None:
        assert current.review_status=="PENDING" and current.publication is None


def test_missing_output_bytes_deny_publication_without_erasing_checks(application):
    state=repair(application)
    state=application.review(SCOPE,REVIEWER,ReviewRequest(expected_version=state.version,note="Sources checked",attested=True))
    # Delete the synthetic test output via the scoped store seam to emulate disk loss.
    with application.blob_store._directory(SCOPE) as directory:
        os.unlink(state.candidate.artifact_id,dir_fd=directory)
    with pytest.raises(DomainConflict,match="missing or corrupt"):
        application.publish(SCOPE,REVIEWER,VersionAction(expected_version=state.version))
    current=application.get_pack(SCOPE,PREPARER)
    assert current.version==state.version and current.checks[0].status=="PASS"


def test_processing_crash_recovery_requires_explicit_stopped_run_recovery(application):
    state=application.recompute(SCOPE,PREPARER,VersionAction(expected_version=1))
    request=application.pending_requests()[0]
    assert application.claim_processing(SCOPE,request["request_id"],"crashed")
    assert not application.claim_processing(SCOPE,request["request_id"],"replacement")
    assert not application.release_processing(SCOPE,request["request_id"],"wrong")
    assert application.release_processing(SCOPE,request["request_id"],"crashed")
    assert application.claim_processing(SCOPE,request["request_id"],"replacement")
    loaded=application.load_processing(SCOPE,request["request_id"])
    result=application.evaluator(loaded["state"],ScopedBlobs(application.blob_store,SCOPE))
    completed=application.complete_processing(SCOPE,request["request_id"],result,"replacement")
    assert completed["applied"] and completed["state"].execution_status=="COMPLETED"


def test_pending_correction_automatically_has_one_processing_request(application):
    state=application.correct(SCOPE,PREPARER,CorrectionRequest(expected_version=1,fact_id="fee",
        value_decimal="10",reason="Correction queues checked computation",source_id="fee-source",idempotency_key="correct-once"))
    assert state.execution_status=="PENDING"
    assert len(application.pending_requests())==1
    again=application.correct(SCOPE,PREPARER,CorrectionRequest(expected_version=1,fact_id="fee",
        value_decimal="10",reason="Correction queues checked computation",source_id="fee-source",idempotency_key="correct-once"))
    assert again.version==state.version and len(application.pending_requests())==1
    coalesced=application.recompute(SCOPE,PREPARER,VersionAction(expected_version=state.version))
    assert coalesced.version==state.version and len(application.pending_requests())==1


def test_publication_commit_serializes_before_concurrent_permission_revocation(application,monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event
    from closegraph.services.repository import PostgreSQLRepository
    state=repair(application)
    state=application.review(SCOPE,REVIEWER,ReviewRequest(expected_version=state.version,note="Sources checked",attested=True))
    at_save,continue_save,revocation_started,revoked=Event(),Event(),Event(),Event()
    original=PostgreSQLRepository.save
    def paused_save(repository,snapshot,expected_version):
        if snapshot.get("publication"):
            at_save.set()
            assert continue_save.wait(5)
        return original(repository,snapshot,expected_version)
    monkeypatch.setattr(PostgreSQLRepository,"save",paused_save)
    def revoke():
        revocation_started.set()
        application.auth.disable("reviewer")
        revoked.set()
    with ThreadPoolExecutor(max_workers=2) as pool:
        publication=pool.submit(application.publish,SCOPE,REVIEWER,VersionAction(expected_version=state.version))
        assert at_save.wait(5)
        revocation=pool.submit(revoke)
        assert revocation_started.wait(5)
        # Revocation cannot report completion before the pending transaction exits.
        assert not revoked.wait(0.1)
        continue_save.set()
        published=publication.result(timeout=5)
        revocation.result(timeout=5)
    assert published.publication is not None and revoked.is_set()
    with pytest.raises(DomainNotFound):
        application.publish(SCOPE,REVIEWER,VersionAction(expected_version=published.version))


def test_revocation_that_finishes_first_denies_publication(application):
    state=repair(application)
    state=application.review(SCOPE,REVIEWER,ReviewRequest(expected_version=state.version,note="Sources checked",attested=True))
    application.auth.disable("reviewer")
    with pytest.raises(DomainNotFound):
        application.publish(SCOPE,REVIEWER,VersionAction(expected_version=state.version))
    assert application.get_pack(SCOPE,PREPARER).publication is None


def test_snapshot_missing_value_and_canonical_parser_provenance(application):
    with application.sessions() as session:
        repository,_=application._service(session,SCOPE)
        state=repository.get(SCOPE.pack_id,SCOPE.fund_id)
    state["facts"][0]["value_decimal"]=None
    state["facts"][0].pop("value_state",None)
    state["source_occurrences"]=[{"occurrence_id":"original-occurrence","source":{"document_version_id":"doc"},
        "raw_content":"Original source text","disposition":"ACCEPTED"}]
    state["extraction_observations"]=[{"observation_id":"native:original-occurrence","occurrence_id":"original-occurrence",
        "parser":"native-csv","parser_version":"1","mode":"NATIVE","settings":{"layout":"declared"},
        "response_id":"doc","confidence":{"status":"NOT_APPLICABLE","reason":"Native extraction has no probability"},
        "candidate":{"raw_value":"1000.00","value_decimal":"1000.00"}}]
    response=application._snapshot(state)
    assert response.facts[0].value_state=="MISSING" and response.facts[0].value_decimal is None
    assert response.extraction_observations[0]["mode"]=="NATIVE"
    assert response.extraction_observations[0]["confidence"]["status"]=="NOT_APPLICABLE"
    assert response.source_occurrences[0]["raw_content"]=="Original source text"
    state["facts"][0]["value_state"]="NOT_APPLICABLE"
    assert application._snapshot(state).facts[0].value_state=="NOT_APPLICABLE"


def test_review_activity_is_scoped_durable_and_does_not_change_financial_revision(application,database):
    from closegraph.contracts import ReviewEventRequest
    from closegraph.api.ports import DomainForbidden
    state=repair(application)
    document_id=state.facts[0].source.document_version_id
    commands=[
        ReviewEventRequest(event_type="SOURCE_OPENED",observed_snapshot_version=state.version,
                           document_version_id=document_id,idempotency_key="source-once"),
        ReviewEventRequest(event_type="VALUE_INSPECTED",observed_snapshot_version=state.version,fact_id="fee"),
        ReviewEventRequest(event_type="READY_ITEM_SAMPLED",observed_snapshot_version=state.version,fact_id="fee"),
    ]
    first=application.record_review_event(SCOPE,REVIEWER,commands[0])
    assert application.record_review_event(SCOPE,REVIEWER,commands[0])==first
    for command in commands[1:]:
        application.record_review_event(SCOPE,REVIEWER,command)
    with pytest.raises(DomainForbidden):
        application.record_review_event(SCOPE,PREPARER,commands[-1])
    with pytest.raises(DomainNotFound):
        application.record_review_event(SCOPE,REVIEWER,commands[0].model_copy(update={"document_version_id":"foreign","idempotency_key":None}))
    with pytest.raises(DomainConflict):
        application.record_review_event(SCOPE,REVIEWER,commands[0].model_copy(update={"fact_id":"fee"}))
    restarted=PostgreSQLPackServices(database,application.blob_store,application.auth,Evaluator())
    work=restarted.review_work(SCOPE,REVIEWER)
    assert work["counts"]=={"SOURCE_OPENED":1,"VALUE_INSPECTED":1,"READY_ITEM_SAMPLED":1}
    assert work["provenance"]=="LOCAL_OBSERVED" and work["elapsed_seconds"]>=0
    assert work["actions"]["CORRECTED"]==1
    assert work["reopened_scope"][0]["impact"]=={"affected":["fee","output"],"independent":[],"unverifiable":False}
    after=application.get_pack(SCOPE,PREPARER)
    assert after.version==state.version and after.candidate==state.candidate
    changed=application.correct(SCOPE,PREPARER,CorrectionRequest(expected_version=state.version,fact_id="fee",
        value_decimal="9",reason="Reopen only actual dependent work",source_id="fee-source"))
    with pytest.raises(DomainConflict,match="current snapshot"):
        application.record_review_event(SCOPE,REVIEWER,commands[-1])
    blocked=commands[-1].model_copy(update={"observed_snapshot_version":changed.version})
    with pytest.raises(DomainConflict,match="checked ready"):
        application.record_review_event(SCOPE,REVIEWER,blocked)
    assert application.review_work(SCOPE,REVIEWER)["total_events"]==3


def test_scoped_independent_output_retains_exact_signoff_and_release(application):
    # Each pack is one output approval boundary. This second investor's output
    # owns distinct evidence identities/context; no shared source/rule dependency.
    from copy import deepcopy
    independent=Scope(tenant_id=SCOPE.tenant_id,fund_id=SCOPE.fund_id,pack_id="independent-investor-output")
    with application.auth._lock:
        for username,account in list(application.auth._accounts.items()):
            application.auth._accounts[username]=account.model_copy(update={"scopes":(SCOPE,independent)})
    with application.sessions() as session:
        repository,_=application._service(session,SCOPE)
        original=repository.get(SCOPE.pack_id,SCOPE.fund_id,version=1)
    state=deepcopy(original)
    state["title"]="Independent investor statement"
    state["rule_version"]="independent-investor-fee-rule-v1"
    state["evidence"]={"independent-"+key:{**value,"entity_id":"independent-investor"}
                       for key,value in original["evidence"].items()}
    state["source_versions"]={key:1 for key in state["evidence"]}
    for fact in state["facts"]:
        fact["entity_id"]="independent-investor"
        fact["source"]={"source_id":"independent-"+fact["source"]["source_id"]}
    for source in state["evidence"].values():
        ScopedBlobs(application.blob_store,independent).put(
            ScopedBlobs(application.blob_store,SCOPE).get(source["content_hash"]))
    ScopedBlobs(application.blob_store,independent).put(
        ScopedBlobs(application.blob_store,SCOPE).get(state["template"]["template_hash"]))
    application.create_pack(independent,PREPARER,state)
    fixed=application.correct(independent,PREPARER,CorrectionRequest(expected_version=1,fact_id="fee",
        value_decimal="10",reason="Independent investor evidence",source_id="independent-fee-source"))
    request=next(r for r in application.pending_requests() if r["scope"]==independent.model_dump())
    assert application.claim_processing(independent,request["request_id"],"independent-run")
    loaded=application.load_processing(independent,request["request_id"])
    result=application.evaluator(loaded["state"],ScopedBlobs(application.blob_store,independent))
    ready=application.complete_processing(independent,request["request_id"],result,"independent-run")["state"]
    approved=application.review(independent,REVIEWER,ReviewRequest(expected_version=ready.version,note="Independent sources reviewed",attested=True))
    released=application.publish(independent,REVIEWER,VersionAction(expected_version=approved.version))
    before_bytes=application.download(independent,REVIEWER,released.candidate.artifact_id)
    primary=application.correct(SCOPE,PREPARER,CorrectionRequest(expected_version=1,fact_id="fee",
        value_decimal="9",reason="Affects only the primary investor statement",source_id="fee-source"))
    assert primary.freshness=="STALE" and primary.impact["affected"]==["fee","output"]
    retained=application.get_pack(independent,REVIEWER)
    assert retained.version==released.version and retained.freshness=="CURRENT"
    assert retained.review==released.review and retained.publication==released.publication
    assert retained.candidate==released.candidate
    after_bytes=application.download(independent,REVIEWER,retained.candidate.artifact_id)
    assert after_bytes==before_bytes and "stale" not in after_bytes.filename


def configuration_change(application, state, kind):
    from closegraph.contracts import ConfigurationChangeRequest
    with application.sessions() as session:
        repository,_=application._service(session,SCOPE)
        internal=repository.get(SCOPE.pack_id,SCOPE.fund_id)
    if kind=="rule":
        changes={"rule_version":"source-fee-v2","rule":{"version":"source-fee-v2","source_id":"fee-source","approval_id":"synthetic-approved"}}
    elif kind=="mapping":
        template=deepcopy(internal["template"])
        changes={"mapping_version":"template-v2","mappings":{"version":"template-v2","template":template},"template":template}
    else:
        changes={"policy_version":"unregistered-policy-v2","policy":{"version":"unregistered-policy-v2","required_check_ids":["fee"]}}
    return ConfigurationChangeRequest(expected_version=state.version,changes=changes,
        reason="Approved server configuration replacement",idempotency_key="configured-"+kind)


@pytest.mark.parametrize("kind",["rule","mapping","policy"])
def test_configured_records_invalidate_signoff_and_enqueue_atomically(application,kind):
    from closegraph.contracts import ConfigurationChangeRequest
    state=repair(application)
    state=application.review(SCOPE,REVIEWER,ReviewRequest(expected_version=state.version,note="Compared sources",attested=True))
    command=configuration_change(application,state,kind)
    changed=application.invalidate(SCOPE,PREPARER,command)
    assert changed.review_status=="PENDING" and changed.candidate is None
    assert changed.execution_status=="PENDING" and changed.freshness=="STALE"
    assert changed.impact["affected"]==sorted([kind,"output"])
    requests=application.pending_requests()
    assert len(requests)==1 and requests[0]["snapshot_version"]==changed.version
    assert application.invalidate(SCOPE,PREPARER,command).version==changed.version
    assert len(application.pending_requests())==1
    loaded=application.load_processing(SCOPE,requests[0]["request_id"])
    assert all(loaded["state"][key]==value for key,value in command.changes.items())
    assert application.claim_processing(SCOPE,requests[0]["request_id"],"new-config")
    result=application.evaluator(loaded["state"],ScopedBlobs(application.blob_store,SCOPE))
    ready=application.complete_processing(SCOPE,requests[0]["request_id"],result,"new-config")["state"]
    with application.sessions() as session:
        repository,_=application._service(session,SCOPE)
        internal=repository.get(SCOPE.pack_id,SCOPE.fund_id)
    assert all(internal[key]==value for key,value in command.changes.items())
    assert ready.review_status=="PENDING"
    if kind=="policy":
        assert ready.routing_status=="NEEDS_REVIEW" and ready.review_status=="PENDING"
    else:
        assert ready.routing_status=="READY_FOR_QUICK_REVIEW"
        approved=application.review(SCOPE,REVIEWER,ReviewRequest(expected_version=ready.version,note="Rechecked new configured version",attested=True))
        released=application.publish(SCOPE,REVIEWER,VersionAction(expected_version=approved.version))
        import json
        manifest=json.loads(application.download(SCOPE,REVIEWER,released.publication["manifest_id"]).content)
        key="rule" if kind=="rule" else "mappings"
        assert manifest["configuration_records"][key]==command.changes[key]
    altered=deepcopy(command.changes)
    record_key={"rule":"rule","mapping":"mappings","policy":"policy"}[kind]
    altered[record_key]={**altered[record_key],"note":"Changed bytes under same version"}
    with pytest.raises(DomainConflict,match="new version"):
        application.invalidate(SCOPE,PREPARER,ConfigurationChangeRequest(expected_version=application.get_pack(SCOPE,PREPARER).version,
            changes=altered,reason="Cannot mutate versioned records"))


@pytest.mark.parametrize("kind",["rule","mapping","policy","source"])
def test_all_relevant_configuration_and_source_writers_race_publication(application,kind):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    state=repair(application)
    state=application.review(SCOPE,REVIEWER,ReviewRequest(expected_version=state.version,note="Sources checked",attested=True))
    if kind=="source":
        command=UploadRequest(expected_version=state.version,idempotency_key="source-race",source_id="fee-source",
            filename="fee.csv",media_type="text/csv",content_base64=base64.b64encode(b"11.00").decode())
        mutate=application.upload
    else:
        command=configuration_change(application,state,kind)
        mutate=application.invalidate
    barrier=Barrier(2)
    def attempt(function,actor,command):
        barrier.wait()
        try:
            return function(SCOPE,actor,command)
        except DomainConflict:
            return None
    with ThreadPoolExecutor(max_workers=2) as pool:
        first=pool.submit(attempt,application.publish,REVIEWER,VersionAction(expected_version=state.version))
        second=pool.submit(attempt,mutate,PREPARER,command)
        outcomes=[first.result(timeout=5),second.result(timeout=5)]
    assert sum(result is not None for result in outcomes)==1
    current=application.get_pack(SCOPE,PREPARER)
    assert current.version==state.version+1
    if outcomes[1] is not None:
        assert current.publication is None and current.candidate is None and current.review_status=="PENDING"
        assert len(application.pending_requests())==1
    else:
        assert current.publication is not None and not application.pending_requests()


def test_configuration_change_rejects_late_worker_result_without_losing_new_request(application):
    pending=application.recompute(SCOPE,PREPARER,VersionAction(expected_version=1))
    request=application.pending_requests()[0]
    loaded=application.load_processing(SCOPE,request["request_id"])
    assert application.claim_processing(SCOPE,request["request_id"],"old-config-run")
    command=configuration_change(application,pending,"rule")
    current=application.invalidate(SCOPE,PREPARER,command)
    old_result=application.evaluator(loaded["state"],ScopedBlobs(application.blob_store,SCOPE))
    late=application.complete_processing(SCOPE,request["request_id"],old_result,"old-config-run")
    assert not late["applied"] and late["state"].version==current.version
    remaining=application.pending_requests()
    assert len(remaining)==1 and remaining[0]["snapshot_version"]==current.version


def test_trusted_configuration_requires_complete_matching_records_and_has_no_http_override(application):
    from pydantic import ValidationError
    from closegraph.contracts import ConfigurationChangeRequest
    for changes in ({"rule_version":"v2"},{"rule_version":"v2","rule":{"version":"v1","approved":True}},
                    {"facts":[]},{"mapping_version":"v2","mappings":{"version":"v2","template":{}}}):
        with pytest.raises(ValidationError):
            ConfigurationChangeRequest(expected_version=1,changes=changes,reason="configured")
