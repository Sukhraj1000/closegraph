from hashlib import sha256
from io import BytesIO
from openpyxl import load_workbook
from closegraph.fixtures import build_fixture, SCOPE, RULE


def test_fixture_is_labelled_scoped_and_has_real_source_anchors(tmp_path):
    files=build_fixture(tmp_path)
    assert SCOPE['fund_id']=='synthetic-fund'
    assert RULE['approved'] is True and RULE['approval_id'].startswith('synthetic-')
    assert b'0012,2026-Q1,GBP,capital,12000000,units' in files['capital.csv'].read_bytes()
    book=load_workbook(BytesIO(files['original.xlsx'].read_bytes()))
    assert book['Pack']['A1'].value.startswith('SYNTHETIC')
    assert book['Pack']['B6'].value==75000
    expected=load_workbook(files['expected.xlsx'])
    assert expected['Pack']['B6'].value==60000 and expected['Pack']['B7'].value==60000
    assert book['Pack']['B2'].value=='0012'
