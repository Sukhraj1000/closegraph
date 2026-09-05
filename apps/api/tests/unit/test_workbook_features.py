from io import BytesIO
from warnings import warn
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from closegraph.outputs import workbook
from closegraph.services.errors import DomainError


def simple_template():
    book = Workbook()
    book.active["A1"] = 1
    book.active["B1"] = "Leave this source cell unchanged"
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def add_unsupported_markup(data, part, *, alternate=False):
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w") as destination:
        for name in source.namelist():
            content = source.read(name)
            if name == part:
                root = ElementTree.fromstring(content)
                if alternate:
                    extra = ElementTree.SubElement(
                        root, "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"
                    )
                    ElementTree.SubElement(extra, "{urn:closegraph:audit}unsupported-content")
                else:
                    extension_list = ElementTree.SubElement(
                        root, "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}extLst"
                    )
                    extension = ElementTree.SubElement(
                        extension_list,
                        "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}ext",
                        {"uri": "{CLOSEGRAPH-UNSUPPORTED-TEMPLATE-FEATURE}"},
                    )
                    ElementTree.SubElement(extension, "{urn:closegraph:audit}must-not-disappear")
                content = ElementTree.tostring(root)
            destination.writestr(name, content)
    return output.getvalue()


@pytest.mark.parametrize("part", [
    "xl/worksheets/sheet1.xml", "xl/workbook.xml", "xl/styles.xml",
])
def test_generation_blocks_extensions_before_a_lossy_loader_can_remove_them(part):
    template = add_unsupported_markup(simple_template(), part)
    contract = workbook.TemplateContract(workbook.digest(template), (("fee", "Sheet", "A1"),))
    with pytest.raises(DomainError, match="Unsupported OOXML markup"):
        workbook.generate_workbook(template, contract, {"fee": "2"})


def test_generation_blocks_unsupported_alternate_content():
    template = add_unsupported_markup(
        simple_template(), "xl/worksheets/sheet1.xml", alternate=True
    )
    contract = workbook.TemplateContract(workbook.digest(template), (("fee", "Sheet", "A1"),))
    with pytest.raises(DomainError, match="Unsupported OOXML markup"):
        workbook.generate_workbook(template, contract, {"fee": "2"})


def test_roundtrip_rejects_unsupported_markup_even_when_only_candidate_is_modified():
    template = simple_template()
    contract = workbook.TemplateContract(workbook.digest(template), (("fee", "Sheet", "A1"),))
    candidate = workbook.generate_workbook(template, contract, {"fee": "2"})
    candidate = add_unsupported_markup(candidate, "xl/worksheets/sheet1.xml")
    with pytest.raises(DomainError, match="Unsupported OOXML markup"):
        workbook.verify_roundtrip(template, candidate, contract, {"fee": "2"})


def test_loader_warnings_block_unrecognized_lossy_features(monkeypatch):
    template = simple_template()
    contract = workbook.TemplateContract(workbook.digest(template), (("fee", "Sheet", "A1"),))
    original_loader = workbook.load_workbook

    def loader_with_warning(*args, **kwargs):
        warn("Unsupported future feature will be removed", UserWarning, stacklevel=2)
        return original_loader(*args, **kwargs)

    monkeypatch.setattr(workbook, "load_workbook", loader_with_warning)
    with pytest.raises(DomainError, match="Unsupported XLSX loader warning"):
        workbook.generate_workbook(template, contract, {"fee": "2"})


def test_xml_screening_rejects_utf16_doctype_before_expanding_entities():
    original = simple_template()
    output = BytesIO()
    with ZipFile(BytesIO(original)) as source, ZipFile(output, "w") as destination:
        for name in source.namelist():
            content = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                text = content.decode()
                text = '<!DOCTYPE worksheet [<!ENTITY audit "do-not-expand">]>' + text
                content = text.encode("utf-16")
            destination.writestr(name, content)
    template = output.getvalue()
    contract = workbook.TemplateContract(workbook.digest(template), (("fee", "Sheet", "A1"),))
    with pytest.raises(DomainError, match="XML entities are unsupported"):
        workbook.generate_workbook(template, contract, {"fee": "2"})


def test_xml_screening_rejects_external_relationship_regardless_of_quote_style():
    original = simple_template()
    output = BytesIO()
    with ZipFile(BytesIO(original)) as source, ZipFile(output, "w") as destination:
        for name in source.namelist():
            content = source.read(name)
            if name == "xl/_rels/workbook.xml.rels":
                content = content.replace(b"<Relationship ", b"<Relationship TargetMode='External' ", 1)
            destination.writestr(name, content)
    template = output.getvalue()
    contract = workbook.TemplateContract(workbook.digest(template), (("fee", "Sheet", "A1"),))
    with pytest.raises(DomainError, match="External references"):
        workbook.generate_workbook(template, contract, {"fee": "2"})


@pytest.mark.parametrize("storage", ["inline", "shared"])
@pytest.mark.parametrize("markup", ["r", "rPh", "phoneticPr"])
def test_unmapped_rich_or_phonetic_text_cannot_be_silently_flattened(storage, markup):
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    tag = lambda name: "{" + namespace + "}" + name
    original = simple_template()
    with ZipFile(BytesIO(original)) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    sheet = ElementTree.fromstring(parts["xl/worksheets/sheet1.xml"])
    cell = next(item for item in sheet.iter(tag("c")) if item.attrib["r"] == "B1")
    text = cell.find(tag("is")).find(tag("t")).text
    for child in list(cell):
        cell.remove(child)
    value = ElementTree.Element(tag("is" if storage == "inline" else "si"))
    if markup == "r":
        first = ElementTree.SubElement(value, tag("r"))
        properties = ElementTree.SubElement(first, tag("rPr"))
        ElementTree.SubElement(properties, tag("b"))
        ElementTree.SubElement(first, tag("t")).text = text[:5]
        second = ElementTree.SubElement(value, tag("r"))
        ElementTree.SubElement(second, tag("t")).text = text[5:]
    else:
        ElementTree.SubElement(value, tag("t")).text = text
        if markup == "rPh":
            phonetic = ElementTree.SubElement(value, tag("rPh"), {"sb": "0", "eb": "5"})
            ElementTree.SubElement(phonetic, tag("t")).text = "reading"
        else:
            ElementTree.SubElement(value, tag("phoneticPr"), {"fontId": "0"})
    if storage == "inline":
        cell.append(value)
    else:
        cell.set("t", "s")
        ElementTree.SubElement(cell, tag("v")).text = "0"
        shared = ElementTree.Element(tag("sst"), {"count": "1", "uniqueCount": "1"})
        shared.append(value)
        parts["xl/sharedStrings.xml"] = ElementTree.tostring(shared)
        types = ElementTree.fromstring(parts["[Content_Types].xml"])
        ElementTree.SubElement(
            types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
            {"PartName": "/xl/sharedStrings.xml",
             "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"},
        )
        parts["[Content_Types].xml"] = ElementTree.tostring(types)
        relations = ElementTree.fromstring(parts["xl/_rels/workbook.xml.rels"])
        ElementTree.SubElement(
            relations, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
            {"Id": "rIdAuditSharedStrings",
             "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
             "Target": "sharedStrings.xml"},
        )
        parts["xl/_rels/workbook.xml.rels"] = ElementTree.tostring(relations)
    parts["xl/worksheets/sheet1.xml"] = ElementTree.tostring(sheet)
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    template = output.getvalue()
    contract = workbook.TemplateContract(workbook.digest(template), (("fee", "Sheet", "A1"),))
    with pytest.raises(DomainError, match="Unsupported rich or phonetic text"):
        workbook.generate_workbook(template, contract, {"fee": "2"})


@pytest.mark.parametrize("feature,attributes", [
    ("customSheetViews", {}),
    ("sheetProtection", {"sheet": "1", "password": "ABCD"}),
    ("headerFooter", {}),
    ("rowBreaks", {"count": "1", "manualBreakCount": "1"}),
])
def test_unrecognized_standard_worksheet_features_block_before_generation(feature, attributes):
    template = simple_template()
    output = BytesIO()
    with ZipFile(BytesIO(template)) as source, ZipFile(output, "w") as destination:
        for name in source.namelist():
            content = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(content)
                tag = lambda name: "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}" + name
                extra = ElementTree.SubElement(root, tag(feature), attributes)
                if feature == "customSheetViews":
                    ElementTree.SubElement(extra, tag("customSheetView"), {
                        "guid": "{01234567-0123-0123-0123-0123456789AB}",
                        "scale": "75", "showGridLines": "0",
                    })
                if feature == "headerFooter":
                    ElementTree.SubElement(extra, tag("oddHeader")).text = "&CMust remain visible"
                if feature == "rowBreaks":
                    ElementTree.SubElement(extra, tag("brk"), {"id": "1", "max": "16383", "man": "1"})
                content = ElementTree.tostring(root)
            destination.writestr(name, content)
    data = output.getvalue()
    contract = workbook.TemplateContract(workbook.digest(data), (("fee", "Sheet", "A1"),))
    with pytest.raises(DomainError, match="Unsupported value-only OOXML"):
        workbook.generate_workbook(data, contract, {"fee": "2"})


def test_unknown_attribute_on_a_supported_element_is_not_silently_ignored():
    data = simple_template()
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w") as destination:
        for name in source.namelist():
            content = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(content)
                root.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheetData").set(
                    "futureSetting", "must-preserve"
                )
                content = ElementTree.tostring(root)
            destination.writestr(name, content)
    data = output.getvalue()
    contract = workbook.TemplateContract(workbook.digest(data), (("fee", "Sheet", "A1"),))
    with pytest.raises(DomainError, match="Unsupported value-only OOXML"):
        workbook.generate_workbook(data, contract, {"fee": "2"})


def layout_template():
    from openpyxl.styles import Font
    from openpyxl.worksheet.views import SheetView
    book = Workbook()
    sheet = book.active
    sheet["A1"] = 1
    sheet["B1"] = "0012"
    sheet["B1"].font = Font(bold=True, color="00336699")
    sheet["A1"].number_format = "#,##0.00"
    sheet.freeze_panes = "B2"
    sheet.column_dimensions["B"].width = 29
    sheet.sheet_view.zoomScale = 90
    sheet.views.sheetView.append(SheetView(workbookViewId=0, zoomScale=75, showGridLines=False))
    book.properties.title = "Original synthetic reporting layout"
    book.properties.keywords = "must remain unchanged"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def test_value_only_writer_preserves_all_other_package_bytes_and_multiple_views():
    from openpyxl import load_workbook
    template = layout_template()
    contract = workbook.TemplateContract(workbook.digest(template), (("fee", "Sheet", "A1"),))
    candidate = workbook.generate_workbook(template, contract, {"fee": "2.50"})
    with ZipFile(BytesIO(template)) as original, ZipFile(BytesIO(candidate)) as revised:
        assert set(original.namelist()) == set(revised.namelist())
        for name in original.namelist():
            if name != "xl/worksheets/sheet1.xml":
                assert original.read(name) == revised.read(name), name
    result = load_workbook(BytesIO(candidate))
    assert result.active["A1"].value == 2.5
    assert result.active["B1"].value == "0012"
    assert result.active["B1"].font.bold
    assert result.active["B1"].font.color.rgb == "00336699"
    assert result.active["A1"].number_format == "#,##0.00"
    assert result.active.column_dimensions["B"].width == 29
    assert result.active.freeze_panes == "B2"
    assert [view.zoomScale for view in result.active.views.sheetView] == [90, 75]
    assert result.properties.title == "Original synthetic reporting layout"
    assert workbook.verify_roundtrip(template, candidate, contract, {"fee": "2.50"})


@pytest.mark.parametrize("part", ["xl/workbook.xml", "docProps/core.xml", "xl/styles.xml",
                                 "xl/worksheets/sheet1.xml"])
def test_verifier_catches_unmapped_metadata_styles_and_nonfirst_view_changes(part):
    template = layout_template()
    contract = workbook.TemplateContract(workbook.digest(template), (("fee", "Sheet", "A1"),))
    candidate = workbook.generate_workbook(template, contract, {"fee": "2.50"})
    output = BytesIO()
    main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(BytesIO(candidate)) as original, ZipFile(output, "w") as revised:
        for name in original.namelist():
            content = original.read(name)
            if name == part:
                root = ElementTree.fromstring(content)
                if part == "xl/workbook.xml":
                    root.find(main + "bookViews/" + main + "workbookView").set("xWindow", "42")
                elif part == "docProps/core.xml":
                    root.find("{http://purl.org/dc/elements/1.1/}title").text = "Tampered title"
                elif part == "xl/styles.xml":
                    root.find(main + "fonts")[1].find(main + "color").set("rgb", "00FF0000")
                else:
                    root.find(main + "sheetViews")[1].set("zoomScale", "50")
                content = ElementTree.tostring(root)
            revised.writestr(name, content)
    with pytest.raises(DomainError, match="Unmapped (package part|worksheet XML) changed"):
        workbook.verify_roundtrip(template, output.getvalue(), contract, {"fee": "2.50"})


@pytest.mark.parametrize("part", [
    "_xmlsignatures/sig1.xml", "xl/vbaProject.bin", "xl/embeddings/data.bin",
    "xl/externalLinks/externalLink1.xml",
])
def test_signatures_macros_embeddings_and_external_parts_are_outside_supported_package(part):
    output = BytesIO()
    with ZipFile(BytesIO(simple_template())) as source, ZipFile(output, "w") as destination:
        for name in source.namelist():
            destination.writestr(name, source.read(name))
        destination.writestr(part, b"unsupported synthetic part")
    template = output.getvalue()
    contract = workbook.TemplateContract(workbook.digest(template), (("fee", "Sheet", "A1"),))
    with pytest.raises(DomainError, match="Unsupported template part"):
        workbook.generate_workbook(template, contract, {"fee": "2"})
