"""Verified, schema-bound CSV/XLSX export of reviewed table data.

Default cells are literal text (including numeric-looking IDs). CSV neutralizes
spreadsheet formula prefixes and records every escape; XLSX emits explicit text
cells. Numeric/date output requires explicit bindings.types. Templates support
plain value/style workbooks only and never overwrite nonblank source content.
"""
from __future__ import annotations

from copy import copy
import csv
from datetime import date, datetime
from decimal import Decimal
import hashlib
from io import BytesIO, StringIO
import json
from pathlib import PurePosixPath
import re
import zipfile
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from openpyxl.styles.numbers import is_date_format

from .transform import _number, _value

MAX_TEMPLATE_BYTES = 30 * 1024 * 1024
MAX_TEMPLATE_EXPANDED = 150 * 1024 * 1024
MAX_EXPORT_ROWS = 200_000
MAX_EXPORT_CELLS = 2_000_000
MAX_TEMPLATE_CELLS = 2_000_000
MAX_XLSX_STRING = 32767
_ALLOWED_PART = re.compile(
    r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:app|core|custom)\.xml|"
    r"xl/(?:workbook\.xml|_rels/workbook\.xml\.rels|styles\.xml|sharedStrings\.xml|"
    r"theme/theme\d+\.xml|worksheets/sheet\d+\.xml))\Z"
)


def _safe_text(value: str | None, *, xlsx: bool) -> str | None:
    if value is None:
        return None
    if ILLEGAL_CHARACTERS_RE.search(value) or "\x00" in value:
        raise ValueError("Output contains unsupported control characters")
    if xlsx and len(value) > MAX_XLSX_STRING:
        raise ValueError("XLSX cell exceeds the 32,767 character limit")
    return value


def _csv_safe(value: str | None) -> tuple[str, bool]:
    text = value or ""
    # Spreadsheet consumers can ignore leading whitespace before interpreting a formula.
    dangerous = text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n"))
    return ("'" + text, True) if dangerous else (text, False)


def _typed(value: str | None, kind: str):
    if kind == "text" or value is None:
        return value
    if kind == "decimal":
        number = _number(value)
        significant = len(number.normalize().as_tuple().digits)
        if significant > 15 or abs(number.adjusted()) > 100:
            raise ValueError("Numeric XLSX value exceeds Excel precision; bind this column as text")
        return float(number)
    if kind == "integer":
        number = _number(value)
        if number != number.to_integral_value() or len(str(abs(int(number)))) > 15:
            raise ValueError("Integer XLSX value exceeds Excel precision or contains a fraction")
        return int(number)
    if kind == "date":
        try:
            parsed = date.fromisoformat(value)
        except ValueError as exc:
            raise ValueError("Date output must use YYYY-MM-DD") from exc
        if parsed < date(1900, 1, 1):
            raise ValueError("XLSX dates before 1900 are unsupported; use text")
        return parsed
    raise ValueError("Output type must be text, decimal, integer or date")


def _load_template(content: bytes):
    if len(content) > MAX_TEMPLATE_BYTES:
        raise ValueError("Template compressed size limit exceeded")
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            names = set()
            expanded = 0
            for entry in archive.infolist():
                name = entry.filename
                if name in names or name.startswith("/") or ".." in PurePosixPath(name).parts or entry.flag_bits & 1:
                    raise ValueError("Template has unsafe, duplicate or encrypted entries")
                names.add(name)
                expanded += entry.file_size
                if expanded > MAX_TEMPLATE_EXPANDED or len(names) > 2048:
                    raise ValueError("Template expanded size or entry limit exceeded")
                if not _ALLOWED_PART.fullmatch(name):
                    raise ValueError("Unsupported template feature or package part: " + name)
            if "xl/workbook.xml" not in names:
                raise ValueError("Template is not a supported XLSX workbook")
            cell_count = 0
            for name in names:
                # Refuse XML entity expansion before either XML reader sees the bytes.
                with archive.open(name) as part:
                    tail = b""
                    while block := part.read(64 * 1024):
                        probe = (tail + block).upper()
                        if b"<!DOCTYPE" in probe or b"<!ENTITY" in probe:
                            raise ValueError("Template XML declarations are unsupported")
                        tail = probe[-16:]
                if name.startswith("xl/worksheets/"):
                    with archive.open(name) as part:
                        try:
                            for _, element in ElementTree.iterparse(part, events=("end",)):
                                if element.tag.rsplit("}", 1)[-1] == "c":
                                    cell_count += 1
                                    if cell_count > MAX_TEMPLATE_CELLS:
                                        raise ValueError("Template cell limit exceeded")
                                element.clear()
                        except ElementTree.ParseError as exc:
                            raise ValueError("Template worksheet XML is malformed") from exc
        workbook = load_workbook(BytesIO(content), data_only=False, keep_links=False)
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ValueError("Template is not a valid XLSX workbook") from exc
    if len(workbook.sheetnames) > 100:
        raise ValueError("Template sheet limit exceeded")
    if sum(len(sheet._cells) for sheet in workbook.worksheets) > MAX_TEMPLATE_CELLS:
        raise ValueError("Template cell limit exceeded")
    for sheet in workbook.worksheets:
        if sheet.tables or sheet.conditional_formatting or sheet.data_validations.count:
            raise ValueError("Template tables, conditional formatting and data validation are unsupported")
        if sheet.protection.sheet:
            raise ValueError("Protected template sheets are unsupported")
        for cell in sheet._cells.values():
            if cell.data_type in ("f", "e") or cell.hyperlink or cell.comment:
                raise ValueError("Template formulas, error cells, hyperlinks and comments are unsupported")
    return workbook


def _snapshot(workbook) -> dict:
    """Semantic content and presentation snapshot; package byte equality is not claimed."""
    sheets = {}
    for sheet in workbook.worksheets:
        cells = {}
        for coord, cell in sheet._cells.items():
            cells[coord] = (cell.value, cell.data_type, copy(cell._style), cell.number_format)
        sheets[sheet.title] = {
            "cells": cells, "merged": set(str(r) for r in sheet.merged_cells.ranges),
            "state": sheet.sheet_state, "freeze_panes": sheet.freeze_panes,
            "rows": {key: (d.height, d.hidden, d.outlineLevel) for key, d in sheet.row_dimensions.items()},
            "columns": {key: (d.min, d.max, d.width, d.hidden, d.outlineLevel) for key, d in sheet.column_dimensions.items()},
            "autofilter": sheet.auto_filter.ref,
        }
    return {"order": list(workbook.sheetnames), "sheets": sheets,
            "names": {key: value.attr_text for key, value in workbook.defined_names.items()}}


def _verify_template(before: dict, workbook, written: set[tuple]) -> None:
    after = _snapshot(workbook)
    if before["order"] != after["order"] or before["names"] != after["names"]:
        raise ValueError("Template sheet order or defined names did not survive readback")
    for title, original in before["sheets"].items():
        current = after["sheets"][title]
        for feature in ("merged", "state", "freeze_panes", "rows", "columns", "autofilter"):
            if original[feature] != current[feature]:
                raise ValueError("Template layout did not survive readback: " + feature)
        for coordinate, cell in original["cells"].items():
            if (title, *coordinate) in written:
                # Existing style must still be retained in intentionally written blank cells.
                actual = current["cells"].get(coordinate)
                if actual is None or cell[2] != actual[2] or cell[3] != actual[3]:
                    raise ValueError("Template target cell formatting did not survive readback")
                continue
            actual = current["cells"].get(coordinate)
            # Empty unstyled cells can legitimately be omitted when serialising XLSX.
            if actual is None and cell[0] is None and not cell[2]:
                continue
            if actual != cell:
                raise ValueError("Unmapped template content did not survive readback")


def export_table(table: dict, format: str = "csv", *, template_bytes: bytes | None = None,
                 bindings: dict | None = None) -> dict:
    """Return bytes, media type, filename and readback verification/source manifest.

    bindings = {sheet?, start_row?, header_row?, columns?:{key:'A'},
                types?:{key:'text'|'decimal'|'integer'|'date'}, labels?:{key:label}}
    Explicit columns order by physical target position. For templates sheet,
    start_row and columns are required; writes are restricted to blank cells.
    """
    if format not in ("csv", "xlsx"):
        raise ValueError("Only CSV and XLSX exports are supported")
    if template_bytes is not None and format != "xlsx":
        raise ValueError("Templates require XLSX export")
    binding = bindings or {}
    if not isinstance(binding, dict) or set(binding) - {"sheet", "start_row", "header_row", "columns", "types", "labels"}:
        raise ValueError("Unsupported export binding fields")
    columns, rows = table.get("columns", []), table.get("rows", [])
    keys = [column["key"] for column in columns]
    if not keys or len(set(keys)) != len(keys) or len(rows) > MAX_EXPORT_ROWS:
        raise ValueError("Export needs distinct columns and at most 200,000 rows")
    mapping = binding.get("columns", {key: get_column_letter(index + 1) for index, key in enumerate(keys)})
    if not isinstance(mapping, dict) or not mapping or set(mapping) - set(keys):
        raise ValueError("Output mapping contains unknown source columns")
    try:
        targets = {key: column_index_from_string(value) for key, value in mapping.items() if isinstance(value, str)}
    except ValueError as exc:
        raise ValueError("Output column targets must be Excel column letters") from exc
    if len(targets) != len(mapping) or len(set(targets.values())) != len(targets) or max(targets.values()) > 16384:
        raise ValueError("Output targets must be distinct valid Excel columns")
    selected = sorted(targets, key=targets.get)
    if len(rows) * len(selected) > MAX_EXPORT_CELLS:
        raise ValueError("Output cell limit exceeded")
    types, labels = binding.get("types", {}), binding.get("labels", {})
    if not isinstance(types, dict) or not isinstance(labels, dict) or (set(types) | set(labels)) - set(selected):
        raise ValueError("Types and labels must refer to selected columns")
    if any(kind not in ("text", "decimal", "integer", "date") for kind in types.values()):
        raise ValueError("Unsupported output type")
    headers = {key: _safe_text(_value(labels.get(key, next(c.get("label", key) for c in columns if c["key"] == key))), xlsx=format == "xlsx") for key in selected}
    manifest = {"table_id": table.get("table_id"), "row_count": len(rows), "columns": selected,
                "cells": [{"row_id": row["row_id"], "columns": {
                    key: row.get("lineage", {}).get(key, [{"table_id": table.get("table_id"), "source_id": table.get("source_id"), "row_id": row["row_id"], "column_key": key, "locator": row.get("locators", {}).get(key)}]) for key in selected}} for row in rows]}
    escapes = []
    if format == "csv":
        stream = StringIO(newline="")
        writer = csv.writer(stream, lineterminator="\r\n")
        expected = []
        header = []
        for key in selected:
            escaped, changed = _csv_safe(headers[key])
            header.append(escaped)
            if changed:
                escapes.append({"row": "header", "column": key})
        writer.writerow(header)
        expected.append(header)
        for row in rows:
            values = []
            for key in selected:
                value = _safe_text(_value(row["values"].get(key)), xlsx=False)
                if types.get(key, "text") != "text":
                    _typed(value, types[key])  # Validate declarations without rounding CSV values.
                escaped, changed = _csv_safe(value)
                values.append(escaped)
                if changed:
                    escapes.append({"row_id": row["row_id"], "column": key})
            writer.writerow(values)
            expected.append(values)
        content = stream.getvalue().encode("utf-8-sig")
        observed = list(csv.reader(StringIO(content.decode("utf-8-sig"), newline="")))
        if observed != expected:
            raise ValueError("CSV readback mismatch")
        verification = {"readback_passed": True, "verified_rows": len(rows), "verified_columns": len(selected),
                        "formula_policy": "dangerous prefixes escaped with apostrophe", "escaped_cells": escapes,
                        "null_policy": "null and empty string both export as empty CSV fields"}
        media_type = "text/csv; charset=utf-8"
    else:
        if template_bytes is not None and not {"sheet", "start_row", "columns"}.issubset(binding):
            raise ValueError("Template requires explicit sheet, start_row and columns bindings")
        workbook = _load_template(template_bytes) if template_bytes is not None else Workbook()
        before = _snapshot(workbook) if template_bytes is not None else None
        title = binding.get("sheet", "Export")
        if not isinstance(title, str) or not title or len(title) > 31 or re.search(r"[\\/*?:\[\]]", title):
            raise ValueError("Invalid worksheet title")
        if template_bytes is not None:
            if title not in workbook.sheetnames:
                raise ValueError("Bound worksheet does not exist in template")
            sheet = workbook[title]
        else:
            sheet = workbook.active
            sheet.title = title
        start_row = binding.get("start_row", 2)
        header_row = binding.get("header_row", 1 if template_bytes is None else None)
        if type(start_row) is not int or start_row < 1 or start_row + len(rows) - 1 > 1048576:
            raise ValueError("Output row binding exceeds worksheet limits")
        if header_row is not None and (type(header_row) is not int or not 1 <= header_row <= 1048576 or start_row <= header_row < start_row + len(rows)):
            raise ValueError("Header row must be outside the output data rows")
        writes = []
        for key in selected:
            if header_row is not None:
                writes.append((header_row, targets[key], headers[key], "text", key))
        for index, row in enumerate(rows, start=start_row):
            for key in selected:
                value = _safe_text(_value(row["values"].get(key)), xlsx=True)
                writes.append((index, targets[key], value, types.get(key, "text"), key))
        written = set()
        expected = []
        for row_number, column_number, value, kind, key in writes:
            cell = sheet.cell(row_number, column_number)
            if any(cell.coordinate in merged for merged in sheet.merged_cells.ranges):
                raise ValueError("Output binding intersects a merged template region")
            if template_bytes is not None and cell.value is not None:
                if row_number != header_row or cell.value != value:
                    raise ValueError("Output binding would overwrite nonblank template content")
                expected.append((row_number, column_number, value, kind))
                continue
            previous_style = copy(cell._style)
            if template_bytes is not None and kind == "date" and value is not None and not is_date_format(cell.number_format):
                raise ValueError("Template date bindings require date-formatted target cells; use text otherwise")
            typed = _typed(value, kind)
            cell.value = typed
            if isinstance(typed, str):
                cell.data_type = "s"  # Even '=...' is literal; never write user text as a formula.
            if template_bytes is not None:
                cell._style = previous_style
            elif kind == "date":
                cell.number_format = "yyyy-mm-dd"
            written.add((title, row_number, column_number))
            expected.append((row_number, column_number, value, kind))
        stream = BytesIO()
        workbook.save(stream)
        content = stream.getvalue()
        observed_book = load_workbook(BytesIO(content), data_only=False)
        observed_sheet = observed_book[title]
        for row_number, column_number, value, kind in expected:
            cell = observed_sheet.cell(row_number, column_number)
            observed = cell.value
            if cell.data_type == "f":
                raise ValueError("Unexpected formula in generated output")
            if value in (None, ""):
                matches = observed in (None, "")
            elif kind in ("decimal", "integer"):
                matches = observed is not None and Decimal(str(observed)) == _number(value)
            elif kind == "date":
                # Blank template cell styles are preserved; openpyxl may return its serial.
                if isinstance(observed, datetime):
                    observed = observed.date()
                if isinstance(observed, (float, int)):
                    from openpyxl.utils.datetime import from_excel
                    converted = from_excel(observed, observed_book.epoch)
                    observed = converted.date() if isinstance(converted, datetime) else converted
                matches = observed == date.fromisoformat(value)
            else:
                matches = isinstance(observed, str) and observed == value
            if not matches:
                raise ValueError(f"XLSX readback mismatch at {get_column_letter(column_number)}{row_number}")
        if before is not None:
            _verify_template(before, observed_book, written)
        verification = {"readback_passed": True, "verified_rows": len(rows), "verified_columns": len(selected),
                        "formula_policy": "all text written as literal strings; formula templates unsupported",
                        "template_preserved": before is not None, "mapped_sheet": title,
                        "null_policy": "null and empty string both read back as empty XLSX cells"}
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    verification["content_sha256"] = hashlib.sha256(content).hexdigest()
    verification["source_manifest"] = manifest
    verification["source_manifest_sha256"] = hashlib.sha256(json.dumps(manifest, sort_keys=True).encode()).hexdigest()
    return {"content": content, "media_type": media_type, "filename": "collection-output." + format,
            "verification": verification}
