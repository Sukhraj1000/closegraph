"""Deliberately bounded value-only XLSX writer; never evaluates Excel formulas."""
import re
import warnings
from copy import copy
from dataclasses import dataclass
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from closegraph.services.corrections import decimal_text
from closegraph.services.errors import DomainError


@dataclass(frozen=True)
class TemplateContract:
    template_hash: str
    bindings: tuple[tuple[str, str, str], ...]  # field, sheet, cell
    version: str = "value-only-v1"


def digest(data):
    return sha256(data).hexdigest()


class _TemplateXML(ElementTree.TreeBuilder):
    def doctype(self, name, pubid, system):
        # Enforce this at the parser boundary, including non-UTF-8 XML.
        raise DomainError("External references or XML entities are unsupported")


# Positive grammar for the interpreted value-only workbook/worksheet profile.
# Other permitted package parts are immutable pass-through assets: the writer
# never reserializes styles, themes, relationships or document properties.
_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def _rule(attributes="", children=""):
    return set(attributes.split()), set(children.split())


_VALUE_XML = {
    "worksheet": _rule(children="sheetPr dimension sheetViews sheetFormatPr cols sheetData autoFilter mergeCells printOptions pageMargins pageSetup"),
    "sheetPr": _rule(children="tabColor outlinePr pageSetUpPr"),
    "tabColor": _rule("rgb indexed auto theme tint"),
    "outlinePr": _rule("applyStyles summaryBelow summaryRight showOutlineSymbols"),
    "pageSetUpPr": _rule("autoPageBreaks fitToPage"),
    "dimension": _rule("ref"),
    "sheetViews": _rule(children="sheetView"),
    "sheetView": _rule("windowProtection showFormulas showGridLines showRowColHeaders showZeros rightToLeft tabSelected showRuler showOutlineSymbols defaultGridColor showWhiteSpace view topLeftCell colorId zoomScale zoomScaleNormal zoomScaleSheetLayoutView zoomScalePageLayoutView workbookViewId", "pane selection"),
    "pane": _rule("xSplit ySplit topLeftCell activePane state"),
    "selection": _rule("pane activeCell activeCellId sqref"),
    "sheetFormatPr": _rule("baseColWidth defaultColWidth defaultRowHeight customHeight zeroHeight thickTop thickBottom outlineLevelRow outlineLevelCol"),
    "cols": _rule(children="col"),
    "col": _rule("min max width style hidden bestFit customWidth outlineLevel collapsed"),
    "sheetData": _rule(children="row"),
    "row": _rule("r spans s customFormat ht hidden customHeight outlineLevel collapsed thickTop thickBot", "c"),
    "c": _rule("r s t", "v is"),
    "v": _rule(),
    "is": _rule(children="t"),
    "t": ({_XML_SPACE}, set()),
    "autoFilter": _rule("ref"),
    "mergeCells": _rule("count", "mergeCell"),
    "mergeCell": _rule("ref"),
    "printOptions": _rule("horizontalCentered verticalCentered headings gridLines gridLinesSet"),
    "pageMargins": _rule("left right top bottom header footer"),
    "pageSetup": _rule("orientation paperSize paperHeight paperWidth scale firstPageNumber fitToHeight fitToWidth pageOrder useFirstPageNumber blackAndWhite draft cellComments errors horizontalDpi verticalDpi copies"),
    "workbook": _rule(children="fileVersion workbookPr workbookProtection bookViews sheets definedNames calcPr"),
    "workbookProtection": _rule(),  # Empty default only; active protection is unsupported.
    "fileVersion": _rule("appName lastEdited lowestEdited rupBuild"),
    "workbookPr": _rule("date1904 showObjects showBorderUnselectedTables filterPrivacy promptedSolutions showInk backupFile saveExternalLinkValues updateLinks hidePivotFieldList publishItems checkCompatibility autoCompressPictures refreshAllConnections defaultThemeVersion"),
    "bookViews": _rule(children="workbookView"),
    "workbookView": _rule("visibility minimized showHorizontalScroll showVerticalScroll showSheetTabs xWindow yWindow windowWidth windowHeight tabRatio firstSheet activeTab autoFilterDateGrouping"),
    "sheets": _rule(children="sheet"),
    "sheet": ({"name", "sheetId", "state", "{" + _REL + "}id"}, set()),
    "definedNames": _rule(),  # Named calculations/print definitions are outside this profile.
    "calcPr": _rule("calcId fullCalcOnLoad forceFullCalc"),
}
_REPEATABLE = {"sheetView", "selection", "col", "row", "c", "mergeCell", "workbookView", "sheet"}


def _validate_value_xml(root, part):
    expected = "workbook" if part == "xl/workbook.xml" else "worksheet"
    if root.tag != "{" + _MAIN + "}" + expected:
        raise DomainError("Unsupported value-only OOXML root in " + part)

    def visit(element):
        name = element.tag.removeprefix("{" + _MAIN + "}")
        if element.tag != "{" + _MAIN + "}" + name or name not in _VALUE_XML:
            raise DomainError("Unsupported value-only OOXML element in " + part)
        attributes, children = _VALUE_XML[name]
        if set(element.attrib) - attributes:
            raise DomainError("Unsupported value-only OOXML attribute in " + part)
        if name not in {"t", "v"} and element.text and element.text.strip():
            raise DomainError("Unsupported value-only OOXML text in " + part)
        seen = set()
        for child in element:
            child_name = child.tag.removeprefix("{" + _MAIN + "}")
            if child_name not in children or (child_name in seen and child_name not in _REPEATABLE):
                raise DomainError("Unsupported value-only OOXML child in " + part)
            if child.tail and child.tail.strip():
                raise DomainError("Unsupported value-only OOXML mixed content in " + part)
            seen.add(child_name)
            visit(child)
        if name == "c" and (len(element) > 1 or element.attrib.get("t", "n") not in
                            {"n", "s", "b", "d", "inlineStr"}):
            raise DomainError("Unsupported value-only OOXML cell in " + part)

    visit(root)


def _xml(data):
    return ElementTree.fromstring(
        data, parser=ElementTree.XMLParser(target=_TemplateXML(insert_comments=True, insert_pis=True))
    )


def _parts(data):
    with ZipFile(BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _worksheet_parts(parts):
    relations = {}
    for item in _xml(parts["xl/_rels/workbook.xml.rels"]):
        identity = item.attrib["Id"]
        if identity in relations:
            raise DomainError("Ambiguous workbook relationships")
        relations[identity] = item.attrib
    sheets = {}
    identities = set()
    workbook = _xml(parts["xl/workbook.xml"])
    for sheet in workbook.find("{" + _MAIN + "}sheets"):
        relation = relations[sheet.attrib["{" + _REL + "}id"]]
        target = relation["Target"]
        path = target.lstrip("/") if target.startswith("/") else "xl/" + target
        if (relation["Type"] != _REL + "/worksheet"
                or not re.fullmatch(r"xl/worksheets/sheet[0-9]+\.xml", path)
                or path not in parts or sheet.attrib["name"] in sheets
                or sheet.attrib["sheetId"] in identities or path in sheets.values()):
            raise DomainError("Unsupported or ambiguous worksheet relationship")
        sheets[sheet.attrib["name"]] = path
        identities.add(sheet.attrib["sheetId"])
    if set(sheets.values()) != {name for name in parts if name.startswith("xl/worksheets/")}:
        raise DomainError("Unreferenced worksheet part")
    return sheets


def _cells(root):
    cells = {}
    rows = set()
    for row in root.findall("{" + _MAIN + "}sheetData/{" + _MAIN + "}row"):
        number = row.attrib.get("r", "")
        if not re.fullmatch(r"[1-9][0-9]*", number) or number in rows:
            raise DomainError("Missing or duplicate worksheet row identity")
        rows.add(number)
        for cell in row:
            coordinate = cell.attrib.get("r", "")
            if (not re.fullmatch(r"[A-Z]{1,3}" + number, coordinate) or coordinate in cells):
                raise DomainError("Missing, duplicate or inconsistent cell identity")
            cells[coordinate] = cell
    return cells


def _xml_signature(element, mapped):
    attributes = dict(element.attrib)
    if element.tag == "{" + _MAIN + "}c" and attributes.get("r") in mapped:
        attributes.pop("t", None)
        return element.tag, tuple(sorted(attributes.items())), None, ()
    text = element.text if element.tag in {"{" + _MAIN + "}t", "{" + _MAIN + "}v"} else None
    return (element.tag, tuple(sorted(attributes.items())), text,
            tuple(_xml_signature(child, mapped) for child in element))


def _verify_package(template, candidate, bindings):
    before, after = _parts(template), _parts(candidate)
    if set(before) != set(after):
        raise DomainError("Output package parts changed")
    sheets = _worksheet_parts(before)
    mapped = {}
    for _, sheet, coordinate in bindings:
        mapped.setdefault(sheets[sheet], set()).add(coordinate)
    for part, original in before.items():
        if part not in mapped:
            if original != after[part]:
                raise DomainError("Unmapped package part changed: " + part)
        elif _xml_signature(_xml(original), mapped[part]) != _xml_signature(_xml(after[part]), mapped[part]):
            raise DomainError("Unmapped worksheet XML changed: " + part)


def _load(data):
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 100 or len({e.filename for e in entries}) != len(entries) or sum(e.file_size for e in entries) > 10_000_000:
                raise DomainError("Unsupported XLSX archive size")
            allowed = re.compile(r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:app|core)\.xml|xl/(?:workbook\.xml|styles\.xml|sharedStrings\.xml|_rels/workbook\.xml\.rels|theme/theme1\.xml|worksheets/sheet[0-9]+\.xml))")
            for entry in entries:
                if not allowed.fullmatch(entry.filename):
                    raise DomainError("Unsupported template part: " + entry.filename)
                content = archive.read(entry)
                if b"<!DOCTYPE" in content.upper() or b"<!ENTITY" in content.upper() or b'TargetMode="External"' in content:
                    raise DomainError("External references or XML entities are unsupported")
                # A parsed workbook cannot prove preservation of XML that its
                # loader ignores. Reject extension/fallback markup before load.
                root = _xml(content)
                if any(not isinstance(element.tag, str) for element in root.iter()):
                    raise DomainError("Unsupported XML comments or processing instructions")
                if any(element.attrib.get("TargetMode") == "External" for element in root.iter()):
                    raise DomainError("External references or XML entities are unsupported")
                if any(element.tag.rsplit("}", 1)[-1] in {"extLst", "AlternateContent"}
                       for element in root.iter()):
                    raise DomainError("Unsupported OOXML markup in " + entry.filename)
                for element in root.iter():
                    if (element.tag.rsplit("}", 1)[-1] in {"is", "si"}
                            and any(child.tag.rsplit("}", 1)[-1] in {"r", "rPh", "phoneticPr"}
                                    for child in element.iter())):
                        raise DomainError("Unsupported rich or phonetic text in " + entry.filename)
                if entry.filename == "xl/workbook.xml" or entry.filename.startswith("xl/worksheets/"):
                    _validate_value_xml(root, entry.filename)
                    if entry.filename.startswith("xl/worksheets/"):
                        _cells(root)
        with warnings.catch_warnings():
            warnings.simplefilter("error")
            workbook = load_workbook(BytesIO(data), data_only=False, keep_links=False)
    except DomainError:
        raise
    except Warning as e:
        raise DomainError("Unsupported XLSX loader warning: " + str(e)) from e
    except Exception as e:
        raise DomainError("Invalid value-only XLSX") from e
    if workbook.defined_names or len(workbook.sheetnames) > 8:
        raise DomainError("Unsupported names or worksheet count")
    for sheet in workbook:
        if sheet.max_row > 10000 or sheet.max_column > 100:
            raise DomainError("Unsupported worksheet bounds")
        if sheet.tables or sheet._charts or sheet._images or sheet.data_validations.count or sheet.conditional_formatting:
            raise DomainError("Unsupported rich template features")
        for row in sheet:
            for cell in row:
                if cell.data_type in {"f", "e"} or cell.comment or cell.hyperlink:
                    raise DomainError("Formula, error, comment or hyperlink unsupported")
    try:
        _worksheet_parts(_parts(data))
    except (KeyError, TypeError, AttributeError) as exc:
        raise DomainError("Invalid worksheet package relationships") from exc
    return workbook


def _signature(workbook, mapped):
    result = []
    for sheet in workbook:
        cells = []
        for row in sheet:
            for cell in row:
                cells.append((cell.coordinate, None if (sheet.title, cell.coordinate) in mapped else cell.value,
                              copy(cell._style), cell.number_format, copy(cell.protection)))
        result.append((sheet.title, sheet.sheet_state, sheet.max_row, sheet.max_column,
                       str(sheet.merged_cells), str(sheet.freeze_panes), repr(sheet.auto_filter),
                       str(sheet.sheet_format), str(sheet.sheet_properties), str(sheet.sheet_view),
                       str(sheet.print_options), str(sheet.page_margins), str(sheet.page_setup),
                       str(sheet.print_area), str(sheet.print_title_rows), str(sheet.print_title_cols),
                       [(k,str(v)) for k,v in sheet.row_dimensions.items()],
                       [(k,str(v)) for k,v in sheet.column_dimensions.items()], cells))
    return result


def verify_roundtrip(template_bytes, candidate_bytes, contract, values):
    if digest(template_bytes) != contract.template_hash or contract.version != "value-only-v1":
        raise DomainError("Template identity/version mismatch")
    original, candidate = _load(template_bytes), _load(candidate_bytes)
    mapped = {(sheet, cell) for _, sheet, cell in contract.bindings}
    fields = [field for field, _, _ in contract.bindings]
    if not fields or len(set(fields)) != len(fields) or len(mapped) != len(fields) or set(fields) != set(values):
        raise DomainError("Output mapping must exactly match the allowlist")
    _verify_package(template_bytes, candidate_bytes, contract.bindings)
    if _signature(original, mapped) != _signature(candidate, mapped):
        raise DomainError("Unmapped cells, styles or layout changed")
    for field, sheet, coordinate in contract.bindings:
        actual = candidate[sheet][coordinate]
        expected = Decimal(decimal_text(values[field]))
        if actual.data_type != "n" or isinstance(actual.value, bool) or actual.value is None or Decimal(str(actual.value)) != expected:
            raise DomainError("Output roundtrip mismatch: " + field)
    return digest(candidate_bytes)


def generate_workbook(template_bytes, contract, values):
    if digest(template_bytes) != contract.template_hash:
        raise DomainError("Template identity mismatch")
    workbook = _load(template_bytes)
    parts = _parts(template_bytes)
    sheet_parts = _worksheet_parts(parts)
    updated = {}
    for field, sheet, coordinate in contract.bindings:
        if field not in values or sheet not in workbook.sheetnames or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", coordinate):
            raise DomainError("Unsupported output binding")
        cell = workbook[sheet][coordinate]
        if isinstance(cell, MergedCell) or cell.row > workbook[sheet].max_row or cell.column > workbook[sheet].max_column:
            raise DomainError("Binding outside supported template")
        value = Decimal(decimal_text(values[field]))
        # Excel cannot preserve more than 15 significant decimal digits reliably.
        if len(value.normalize().as_tuple().digits) > 15:
            raise DomainError("Value exceeds supported Excel precision")
        part = sheet_parts[sheet]
        root = updated.setdefault(part, _xml(parts[part]))
        xml_cell = _cells(root).get(coordinate)
        if xml_cell is None:
            raise DomainError("Mapped cell must exist explicitly in the template XML")
        for child in list(xml_cell):
            xml_cell.remove(child)
        xml_cell.set("t", "n")
        ElementTree.SubElement(xml_cell, "{" + _MAIN + "}v").text = format(value, "f")
    for part, root in updated.items():
        parts[part] = ElementTree.tostring(root, encoding="utf-8")
    stream = BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        for part, content in parts.items():
            archive.writestr(part, content)
    candidate = stream.getvalue()
    verify_roundtrip(template_bytes, candidate, contract, values)
    return candidate
